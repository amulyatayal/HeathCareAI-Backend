from pathlib import Path
import openpyxl
import pytest
from scripts.recipe.apply_recipe_review import read_review_rows, ReviewError

HEADERS = ["Recipe ID", "Title", "Section", "Diet tag", "Allergens listed in book",
           "kcal / serving", "Protein g", "Book description", "Status", "Servings",
           "Prep (min)", "Cook (min)", "Ingredients", "Method",
           "Allergens in the AI ingredients", "Flags for the reviewer", "Decision",
           "Corrections (write the corrected text)", "Comments", "Reviewed by", "Date"]

def make_book(tmp_path: Path, rows: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "Read me"
    ws = wb.create_sheet("AI-drafted (to review)")
    ws.append(["group header row"])
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    wb.create_sheet("Complete in book")
    p = tmp_path / "review.xlsx"
    wb.save(p)
    return p

def row(**kw) -> dict:
    base = {"Recipe ID": "recipe-100", "Title": "Test Soup", "Servings": 4,
            "Prep (min)": 10, "Cook (min)": 20,
            "Ingredients": "• 1 onion\n• 500ml vegetable stock",
            "Method": "1. Chop the onion.\n2. Simmer in stock for 20 minutes.",
            "Allergens in the AI ingredients": "celery"}
    base.update(kw)
    return base

def test_reads_decision_case_insensitively(tmp_path):
    p = make_book(tmp_path, [row(Decision=" approve ", **{"Reviewed by": "J. Patel RD", "Date": "2026-11-01"})])
    [r] = read_review_rows(p)
    assert r.decision == "Approve"
    assert r.ingredients == ["1 onion", "500ml vegetable stock"]
    assert r.instructions == ["Chop the onion.", "Simmer in stock for 20 minutes."]

def test_blank_decision_is_pending(tmp_path):
    [r] = read_review_rows(make_book(tmp_path, [row()]))
    assert r.decision is None

def test_unknown_decision_aborts_with_row_id(tmp_path):
    with pytest.raises(ReviewError, match="recipe-100"):
        read_review_rows(make_book(tmp_path, [row(Decision="Maybe")]))

def test_approved_without_reviewer_name_aborts(tmp_path):
    with pytest.raises(ReviewError, match="Reviewed by"):
        read_review_rows(make_book(tmp_path, [row(Decision="Approve")]))


from scripts.recipe.apply_recipe_review import (  # noqa: E402
    to_app_allergens, allergen_info, merge, ReviewRow,
)


def rr(**kw) -> ReviewRow:
    base = dict(recipe_id="recipe-100", decision="Approve", servings=4, prep_time_mins=10, cook_time_mins=20,
                ingredients=["1 onion"], instructions=["Chop."], allergens_uk=["celery"], allergens_blank=False, flags="",
                corrections="", reviewed_by="J. Patel RD", review_date="2026-11-01")
    base.update(kw)
    return ReviewRow(**base)


def doc_with(rid="recipe-100") -> dict:
    return {"version": "1.0", "generated_at": "x", "generator": "parse_recipes.py", "source": "s", "sections": [],
            "statistics": {"total_recipes": 1, "complete_recipes": 0, "metadata_only_recipes": 1, "with_images": 0, "categories_covered": []},
            "recipes": [{"recipe_id": rid, "title": "Test Soup", "meal_type": "soup", "dietary_tags": ["vegan"],
                         "allergens": [], "allergen_info": "None", "ingredients": [], "instructions": [],
                         "has_full_recipe": False, "nutrition": {}, "servings": None,
                         "prep_time_mins": None, "cook_time_mins": None, "image": None}]}


def test_shellfish_mapping():
    assert to_app_allergens(["crustaceans", "molluscs", "peanuts", "celery"]) == ["celery", "nuts", "shellfish"]


def test_allergen_info_keeps_uk_detail_and_sulphites():
    s = allergen_info(["crustaceans", "gluten"], flags="POSSIBLE SULPHITES: dried apricots")
    assert s == "Contains crustaceans (shellfish), gluten. May contain sulphites (check labels)."


def test_approved_row_completes_recipe():
    new, counts = merge(doc_with(), [rr()])
    r = new["recipes"][0]
    assert r["ingredients"] == ["1 onion"] and r["has_full_recipe"] is True
    assert r["review"] == {"decision": "Approve", "reviewed_by": "J. Patel RD", "date": "2026-11-01",
                           "book_allergens": []}
    assert new["statistics"]["complete_recipes"] == 1 and counts["merged"] == 1


def test_pending_and_rejected_rows_unchanged():
    d = doc_with()
    new, counts = merge(d, [rr(decision=None)])
    assert new["recipes"][0]["ingredients"] == [] and counts["pending"] == 1
    new, counts = merge(d, [rr(decision="Reject")])
    assert new["recipes"][0]["ingredients"] == [] and counts["rejected"] == 1


def test_unapplied_edits_refused():
    with pytest.raises(ReviewError, match="not applied"):
        merge(doc_with(), [rr(decision="Approve with edits", corrections="Ingredients: 2 onions")])


def test_idempotent():
    once, _ = merge(doc_with(), [rr()])
    twice, _ = merge(once, [rr()])
    assert once["recipes"] == twice["recipes"] and once["statistics"] == twice["statistics"]


def test_shellfish_filter_end_to_end():
    from services.recipe_service import _normalise_allergen
    new, _ = merge(doc_with(), [rr(allergens_uk=["molluscs"])])
    assert "shellfish" in {_normalise_allergen(a) for a in new["recipes"][0]["allergens"]}


def test_synonyms_map_to_app_vocabulary():
    assert to_app_allergens(["soya", "milk", "eggs"]) == ["dairy", "egg", "soy"]
    assert to_app_allergens(["cereals containing gluten", "tree nuts", "shell fish"]) == ["gluten", "nuts", "shellfish"]


def test_unknown_allergen_on_approved_row_fails_closed():
    with pytest.raises(ReviewError, match=r"recipe-100.*kiwi"):
        merge(doc_with(), [rr(allergens_uk=["kiwi"])])


def test_unknown_allergen_on_pending_row_ignored():
    _, counts = merge(doc_with(), [rr(decision=None, allergens_uk=["kiwi"])])
    assert counts["pending"] == 1


def test_known_vocabulary_maps_into_app_vocabulary():
    from scripts.recipe.apply_recipe_review import _KNOWN_UK
    app = {"dairy", "gluten", "nuts", "soy", "egg", "fish", "shellfish", "sesame", "celery", "mustard", "lupin", "sulphites"}
    assert set(to_app_allergens(sorted(_KNOWN_UK | {"mustard"}))) <= app


def test_book_allergens_are_never_dropped_and_merge_is_idempotent():
    d = doc_with()
    d["recipes"][0]["allergens"] = ["nuts"]
    once, _ = merge(d, [rr(allergens_uk=["celery"])])
    r = once["recipes"][0]
    assert r["allergens"] == ["celery", "nuts"]
    assert "nuts" in r["allergen_info"] and "The recipe book also lists: nuts." in r["allergen_info"]
    assert r["review"]["book_allergens"] == ["nuts"]
    twice, _ = merge(once, [rr(allergens_uk=["celery"])])
    assert once == twice


def test_blank_allergen_cell_on_approved_row_raises():
    with pytest.raises(ReviewError, match=r"recipe-100: allergen cell is blank; write 'none'"):
        merge(doc_with(), [rr(allergens_uk=[], allergens_blank=True)])


def test_explicit_none_allergens_merges_but_keeps_book_allergens():
    d = doc_with()
    d["recipes"][0]["allergens"] = ["dairy"]
    new, counts = merge(d, [rr(allergens_uk=[], allergens_blank=False)])
    assert counts["merged"] == 1 and new["recipes"][0]["allergens"] == ["dairy"]


def test_reader_flags_blank_allergen_cell_and_none(tmp_path):
    [blank] = read_review_rows(make_book(tmp_path, [row(**{"Allergens in the AI ingredients": "  "})]))
    assert blank.allergens_blank is True
    [none] = read_review_rows(make_book(tmp_path, [row(**{"Allergens in the AI ingredients": "none"})]))
    assert none.allergens_blank is False and none.allergens_uk == []


@pytest.mark.parametrize("decision", [None, "Reject"])
def test_withdrawn_approval_fails_closed(decision):
    merged, _ = merge(doc_with(), [rr()])
    with pytest.raises(ReviewError, match="recipe-100: approval withdrawn but recipe is already merged"):
        merge(merged, [rr(decision=decision)])


def test_complete_book_recipe_is_not_overwritten():
    d = doc_with()
    d["recipes"][0]["has_full_recipe"] = True
    with pytest.raises(ReviewError, match="recipe-100: recipe is complete in the book; refusing to overwrite"):
        merge(d, [rr()])


def test_approved_row_without_date_aborts(tmp_path):
    with pytest.raises(ReviewError, match="Date"):
        read_review_rows(make_book(tmp_path, [row(Decision="Approve", **{"Reviewed by": "J. Patel RD"})]))
    with pytest.raises(ReviewError, match="Date"):
        merge(doc_with(), [rr(review_date="")])


def test_duplicate_recipe_ids_raise():
    with pytest.raises(ReviewError, match="recipe-100"):
        merge(doc_with(), [rr(), rr(decision=None)])


def test_approve_with_edits_marked_applied_merges():
    new, counts = merge(doc_with(), [rr(decision="Approve with edits", corrections="[applied] swapped stock")])
    assert counts["merged"] == 1 and new["recipes"][0]["review"]["decision"] == "Approve with edits"


def test_recipe_id_not_found_raises():
    with pytest.raises(ReviewError, match="recipe-999: not found"):
        merge(doc_with(), [rr(recipe_id="recipe-999")])


def _main_env(tmp_path, monkeypatch, doc, workbook_rows):
    import json
    import sys
    from scripts.recipe import apply_recipe_review as mod
    rec, schema = tmp_path / "recipes_full.json", tmp_path / "schema.json"
    rec.write_text(json.dumps(doc))
    schema.write_text(mod.SCHEMA.read_text())
    book = make_book(tmp_path, workbook_rows)
    monkeypatch.setattr(mod, "RECIPES", rec)
    monkeypatch.setattr(mod, "SCHEMA", schema)
    return mod, rec, book


def test_main_dry_run_does_not_write(tmp_path, monkeypatch, capsys):
    import sys
    rows = [row(Decision="Approve", **{"Reviewed by": "J. Patel RD", "Date": "2026-11-01"})]
    mod, rec, book = _main_env(tmp_path, monkeypatch, doc_with(), rows)
    before = rec.read_bytes()
    monkeypatch.setattr(sys, "argv", ["x", "--workbook", str(book), "--dry-run"])
    mod.main()
    assert rec.read_bytes() == before
    assert "merged=1" in capsys.readouterr().out


def test_main_real_write_prints_reindex_hint(tmp_path, monkeypatch, capsys):
    import json
    import sys
    rows = [row(Decision="Approve", **{"Reviewed by": "J. Patel RD", "Date": "2026-11-01"})]
    mod, rec, book = _main_env(tmp_path, monkeypatch, doc_with(), rows)
    monkeypatch.setattr(sys, "argv", ["x", "--workbook", str(book)])
    mod.main()
    assert json.loads(rec.read_text())["recipes"][0]["has_full_recipe"] is True
    assert "ingest_recipes.py --recreate" in capsys.readouterr().out


def test_shellfish_allergy_filter_excludes_merged_molluscs_recipe():
    from services.recipe_service import RecipeService
    new, _ = merge(doc_with(), [rr(allergens_uk=["molluscs"])])
    safe = dict(new["recipes"][0], recipe_id="recipe-101", allergens=[])
    svc = RecipeService()
    svc._catalogue = [new["recipes"][0], safe]
    got = svc._filter_locally("vegan", ["Shellfish"], 10)
    assert [r["recipe_id"] for r in got] == ["recipe-101"]
